from fastapi import FastAPI, APIRouter, HTTPException, Body
from fastapi.responses import JSONResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Any, Dict
import uuid
from datetime import datetime, timezone
import sqlite3
import psycopg2
import pymysql
try:
    import pyodbc
    PYODBC_AVAILABLE = True
except ImportError:
    PYODBC_AVAILABLE = False
from cryptography.fernet import Fernet
import base64
import hashlib
import json
from transformers import pipeline
import torch
import csv
import io
import sqlparse
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# Initialize local LLM for debugging (lightweight model)
try:
    # Using a small text generation model for SQL debugging
    sql_debugger = pipeline("text-generation", model="distilgpt2", device=-1)  # CPU only
except Exception as e:
    logging.warning(f"Could not load LLM model: {e}")
    sql_debugger = None

# Models
class QueryExecute(BaseModel):
    database_type: str  # sqlite, mysql, postgresql, sqlserver
    query: str
    connection_params: Optional[Dict[str, Any]] = None

class QueryResult(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    success: bool
    columns: Optional[List[str]] = None
    rows: Optional[List[List[Any]]] = None
    affected_rows: Optional[int] = None
    error: Optional[str] = None
    execution_time: Optional[float] = None

class SavedQuery(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    query: str
    database_type: str
    encrypted: bool = False
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class SavedQueryCreate(BaseModel):
    name: str
    query: str
    database_type: str
    encrypted: bool = False
    password: Optional[str] = None

class QueryHistory(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    query: str
    database_type: str
    success: bool
    timestamp: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class DebugRequest(BaseModel):
    query: str
    error_message: Optional[str] = None

class DebugResponse(BaseModel):
    suggestions: List[str]
    fixed_query: Optional[str] = None

class AppSettings(BaseModel):
    model_config = ConfigDict(extra="ignore")
    
    user_id: str = "default"
    debug_enabled: bool = True
    auto_save_history: bool = True
    syntax_highlighting: bool = True
    max_history_items: int = 100

# Helper functions
def encrypt_data(data: str, password: str) -> str:
    """Encrypt data using password-based encryption"""
    key = base64.urlsafe_b64encode(hashlib.sha256(password.encode()).digest())
    fernet = Fernet(key)
    encrypted = fernet.encrypt(data.encode())
    return base64.b64encode(encrypted).decode()

def decrypt_data(encrypted_data: str, password: str) -> str:
    """Decrypt data using password-based encryption"""
    try:
        key = base64.urlsafe_b64encode(hashlib.sha256(password.encode()).digest())
        fernet = Fernet(key)
        decrypted = fernet.decrypt(base64.b64decode(encrypted_data.encode()))
        return decrypted.decode()
    except Exception as e:
        raise HTTPException(status_code=400, detail="Invalid password or corrupted data")

def execute_sqlite(query: str, params: Optional[Dict] = None) -> QueryResult:
    """Execute SQLite query"""
    import time
    start = time.time()
    
    try:
        db_path = params.get('database', ':memory:') if params else ':memory:'
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        cursor.execute(query)
        
        if query.strip().upper().startswith(('SELECT', 'PRAGMA', 'EXPLAIN')):
            columns = [desc[0] for desc in cursor.description] if cursor.description else []
            rows = cursor.fetchall()
            result = QueryResult(
                success=True,
                columns=columns,
                rows=rows,
                execution_time=time.time() - start
            )
        else:
            conn.commit()
            result = QueryResult(
                success=True,
                affected_rows=cursor.rowcount,
                execution_time=time.time() - start
            )
        
        conn.close()
        return result
    except Exception as e:
        return QueryResult(success=False, error=str(e), execution_time=time.time() - start)

def execute_mysql(query: str, params: Optional[Dict] = None) -> QueryResult:
    """Execute MySQL query"""
    import time
    start = time.time()
    
    try:
        if not params:
            raise ValueError("MySQL connection parameters required")
        
        conn = pymysql.connect(
            host=params.get('host', 'localhost'),
            port=params.get('port', 3306),
            user=params.get('user', 'root'),
            password=params.get('password', ''),
            database=params.get('database', ''),
            charset='utf8mb4'
        )
        cursor = conn.cursor()
        
        cursor.execute(query)
        
        if query.strip().upper().startswith(('SELECT', 'SHOW', 'DESCRIBE', 'EXPLAIN')):
            columns = [desc[0] for desc in cursor.description] if cursor.description else []
            rows = cursor.fetchall()
            result = QueryResult(
                success=True,
                columns=columns,
                rows=rows,
                execution_time=time.time() - start
            )
        else:
            conn.commit()
            result = QueryResult(
                success=True,
                affected_rows=cursor.rowcount,
                execution_time=time.time() - start
            )
        
        conn.close()
        return result
    except Exception as e:
        return QueryResult(success=False, error=str(e), execution_time=time.time() - start)

def execute_postgresql(query: str, params: Optional[Dict] = None) -> QueryResult:
    """Execute PostgreSQL query"""
    import time
    start = time.time()
    
    try:
        if not params:
            raise ValueError("PostgreSQL connection parameters required")
        
        conn = psycopg2.connect(
            host=params.get('host', 'localhost'),
            port=params.get('port', 5432),
            user=params.get('user', 'postgres'),
            password=params.get('password', ''),
            database=params.get('database', 'postgres')
        )
        cursor = conn.cursor()
        
        cursor.execute(query)
        
        if query.strip().upper().startswith(('SELECT', 'SHOW', 'EXPLAIN')):
            columns = [desc[0] for desc in cursor.description] if cursor.description else []
            rows = cursor.fetchall()
            result = QueryResult(
                success=True,
                columns=columns,
                rows=rows,
                execution_time=time.time() - start
            )
        else:
            conn.commit()
            result = QueryResult(
                success=True,
                affected_rows=cursor.rowcount,
                execution_time=time.time() - start
            )
        
        conn.close()
        return result
    except Exception as e:
        return QueryResult(success=False, error=str(e), execution_time=time.time() - start)

def execute_sqlserver(query: str, params: Optional[Dict] = None) -> QueryResult:
    """Execute SQL Server query"""
    import time
    start = time.time()
    
    if not PYODBC_AVAILABLE:
        return QueryResult(success=False, error="SQL Server support not available (pyodbc not installed)", execution_time=0)
    
    try:
        if not params:
            raise ValueError("SQL Server connection parameters required")
        
        conn_str = (
            f"DRIVER={{ODBC Driver 17 for SQL Server}};"
            f"SERVER={params.get('host', 'localhost')},{params.get('port', 1433)};"
            f"DATABASE={params.get('database', 'master')};"
            f"UID={params.get('user', 'sa')};"
            f"PWD={params.get('password', '')}"
        )
        conn = pyodbc.connect(conn_str)
        cursor = conn.cursor()
        
        cursor.execute(query)
        
        if query.strip().upper().startswith(('SELECT', 'SHOW', 'EXEC sp_')):
            columns = [desc[0] for desc in cursor.description] if cursor.description else []
            rows = cursor.fetchall()
            result = QueryResult(
                success=True,
                columns=columns,
                rows=[list(row) for row in rows],
                execution_time=time.time() - start
            )
        else:
            conn.commit()
            result = QueryResult(
                success=True,
                affected_rows=cursor.rowcount,
                execution_time=time.time() - start
            )
        
        conn.close()
        return result
    except Exception as e:
        return QueryResult(success=False, error=str(e), execution_time=time.time() - start)

def debug_sql_query(query: str, error_message: Optional[str] = None) -> DebugResponse:
    """Use local LLM to debug SQL query"""
    suggestions = []
    
    # Rule-based debugging (always works)
    if error_message:
        error_lower = error_message.lower()
        
        if "syntax error" in error_lower:
            suggestions.append("Check for missing or extra commas, parentheses, or semicolons")
            suggestions.append("Verify that all SQL keywords are spelled correctly")
        
        if "table" in error_lower and ("not exist" in error_lower or "not found" in error_lower):
            suggestions.append("The table name may not exist. Check table name spelling and case sensitivity")
            suggestions.append("Ensure you're connected to the correct database")
        
        if "column" in error_lower and ("not exist" in error_lower or "not found" in error_lower):
            suggestions.append("The column name may not exist. Verify column names in your table")
        
        if "permission" in error_lower or "access denied" in error_lower:
            suggestions.append("Check database user permissions")
        
        if "connection" in error_lower or "refused" in error_lower:
            suggestions.append("Verify database connection parameters (host, port, credentials)")
    
    # Add general SQL suggestions
    query_upper = query.upper()
    if "SELECT" in query_upper and "FROM" not in query_upper:
        suggestions.append("SELECT statement requires a FROM clause")
    
    if "INSERT INTO" in query_upper and "VALUES" not in query_upper:
        suggestions.append("INSERT statement typically requires VALUES clause")
    
    # Try LLM if available (optional enhancement)
    if sql_debugger and len(suggestions) < 3:
        try:
            prompt = f"SQL Query Debug:\nQuery: {query}\nError: {error_message}\nSuggestion:"
            llm_response = sql_debugger(prompt, max_length=100, num_return_sequences=1)
            if llm_response and len(llm_response) > 0:
                llm_text = llm_response[0].get('generated_text', '')
                if 'Suggestion:' in llm_text:
                    suggestion = llm_text.split('Suggestion:')[-1].strip()
                    if suggestion and len(suggestion) > 10:
                        suggestions.append(suggestion[:200])
        except:
            pass
    
    if not suggestions:
        suggestions.append("Review SQL syntax and database schema")
        suggestions.append("Check for typos in table or column names")
    
    return DebugResponse(suggestions=suggestions, fixed_query=None)

# API Routes
@api_router.get("/")
async def root():
    return {"message": "SQL Studio API", "status": "running"}

@api_router.post("/query/execute", response_model=QueryResult)
async def execute_query(request: QueryExecute):
    """Execute SQL query on specified database"""
    db_type = request.database_type.lower()
    
    if db_type == "sqlite":
        result = execute_sqlite(request.query, request.connection_params)
    elif db_type == "mysql":
        result = execute_mysql(request.query, request.connection_params)
    elif db_type == "postgresql":
        result = execute_postgresql(request.query, request.connection_params)
    elif db_type == "sqlserver":
        result = execute_sqlserver(request.query, request.connection_params)
    else:
        raise HTTPException(status_code=400, detail="Unsupported database type")
    
    # Save to history
    history_entry = QueryHistory(
        query=request.query,
        database_type=db_type,
        success=result.success
    )
    history_dict = history_entry.model_dump()
    history_dict['timestamp'] = history_dict['timestamp'].isoformat()
    await db.query_history.insert_one(history_dict)
    
    return result

@api_router.get("/query/history", response_model=List[QueryHistory])
async def get_query_history(limit: int = 50):
    """Get query execution history"""
    history = await db.query_history.find({}, {"_id": 0}).sort("timestamp", -1).limit(limit).to_list(limit)
    
    for item in history:
        if isinstance(item['timestamp'], str):
            item['timestamp'] = datetime.fromisoformat(item['timestamp'])
    
    return history

@api_router.delete("/query/history")
async def clear_query_history():
    """Clear all query history"""
    result = await db.query_history.delete_many({})
    return {"deleted_count": result.deleted_count}

@api_router.post("/query/save", response_model=SavedQuery)
async def save_query(request: SavedQueryCreate):
    """Save a query (with optional encryption)"""
    query_text = request.query
    
    if request.encrypted and request.password:
        query_text = encrypt_data(request.query, request.password)
    
    saved_query = SavedQuery(
        name=request.name,
        query=query_text,
        database_type=request.database_type,
        encrypted=request.encrypted
    )
    
    query_dict = saved_query.model_dump()
    query_dict['created_at'] = query_dict['created_at'].isoformat()
    await db.saved_queries.insert_one(query_dict)
    
    return saved_query

@api_router.get("/query/saved", response_model=List[SavedQuery])
async def get_saved_queries():
    """Get all saved queries"""
    queries = await db.saved_queries.find({}, {"_id": 0}).sort("created_at", -1).to_list(100)
    
    for query in queries:
        if isinstance(query['created_at'], str):
            query['created_at'] = datetime.fromisoformat(query['created_at'])
    
    return queries

@api_router.post("/query/decrypt")
async def decrypt_query(query_id: str = Body(..., embed=True), password: str = Body(..., embed=True)):
    """Decrypt a saved query"""
    query = await db.saved_queries.find_one({"id": query_id}, {"_id": 0})
    
    if not query:
        raise HTTPException(status_code=404, detail="Query not found")
    
    if not query.get('encrypted'):
        raise HTTPException(status_code=400, detail="Query is not encrypted")
    
    decrypted_text = decrypt_data(query['query'], password)
    return {"query": decrypted_text}

@api_router.delete("/query/saved/{query_id}")
async def delete_saved_query(query_id: str):
    """Delete a saved query"""
    result = await db.saved_queries.delete_one({"id": query_id})
    
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Query not found")
    
    return {"message": "Query deleted successfully"}

@api_router.post("/debug", response_model=DebugResponse)
async def debug_query(request: DebugRequest):
    """Debug SQL query and provide suggestions"""
    return debug_sql_query(request.query, request.error_message)

@api_router.get("/settings", response_model=AppSettings)
async def get_settings():
    """Get app settings"""
    settings = await db.app_settings.find_one({"user_id": "default"}, {"_id": 0})
    
    if not settings:
        settings = AppSettings().model_dump()
        await db.app_settings.insert_one(settings)
    
    return AppSettings(**settings)

@api_router.put("/settings", response_model=AppSettings)
async def update_settings(settings: AppSettings):
    """Update app settings"""
    settings_dict = settings.model_dump()
    
    await db.app_settings.update_one(
        {"user_id": "default"},
        {"$set": settings_dict},
        upsert=True
    )
    
    return settings

@api_router.get("/sample-database")
async def create_sample_database():
    """Create a sample SQLite database for testing"""
    try:
        db_path = "/tmp/sample_database.db"
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        # Create sample tables
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT UNIQUE NOT NULL,
                age INTEGER,
                department TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                product TEXT NOT NULL,
                amount DECIMAL(10, 2),
                status TEXT DEFAULT 'pending',
                order_date TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                category TEXT,
                price DECIMAL(10, 2),
                stock INTEGER DEFAULT 0
            )
        """)
        
        # Insert sample data
        cursor.execute("DELETE FROM orders")
        cursor.execute("DELETE FROM users")
        cursor.execute("DELETE FROM products")
        
        cursor.executemany(
            "INSERT INTO users (name, email, age, department) VALUES (?, ?, ?, ?)",
            [
                ("John Doe", "john@example.com", 30, "Engineering"),
                ("Jane Smith", "jane@example.com", 25, "Marketing"),
                ("Bob Johnson", "bob@example.com", 35, "Sales"),
                ("Alice Brown", "alice@example.com", 28, "Engineering"),
                ("Charlie Davis", "charlie@example.com", 32, "HR")
            ]
        )
        
        cursor.executemany(
            "INSERT INTO products (name, category, price, stock) VALUES (?, ?, ?, ?)",
            [
                ("Laptop", "Electronics", 999.99, 50),
                ("Mouse", "Electronics", 29.99, 200),
                ("Keyboard", "Electronics", 79.99, 150),
                ("Monitor", "Electronics", 299.99, 75),
                ("Headphones", "Electronics", 149.99, 100)
            ]
        )
        
        cursor.executemany(
            "INSERT INTO orders (user_id, product, amount, status) VALUES (?, ?, ?, ?)",
            [
                (1, "Laptop", 999.99, "completed"),
                (1, "Mouse", 29.99, "completed"),
                (2, "Keyboard", 79.99, "pending"),
                (3, "Monitor", 299.99, "completed"),
                (4, "Headphones", 149.99, "shipped")
            ]
        )
        
        conn.commit()
        conn.close()
        
        return {
            "success": True,
            "database_path": db_path,
            "message": "Sample database created successfully",
            "tables": ["users", "orders", "products"],
            "sample_queries": [
                "SELECT * FROM users;",
                "SELECT * FROM orders;",
                "SELECT * FROM products;",
                "SELECT u.name, o.product, o.amount FROM users u JOIN orders o ON u.id = o.user_id;",
                "SELECT COUNT(*) as total_users FROM users;",
                "SELECT SUM(amount) as total_revenue FROM orders WHERE status = 'completed';"
            ]
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# Professional Features APIs

class ExportRequest(BaseModel):
    format: str  # csv, json, excel, sql
    data: Dict[str, Any]
    filename: Optional[str] = "export"

class SchemaRequest(BaseModel):
    database_type: str
    connection_params: Optional[Dict[str, Any]] = None

class FormatQueryRequest(BaseModel):
    query: str

class ImportDataRequest(BaseModel):
    database_type: str
    table_name: str
    data: List[Dict[str, Any]]
    connection_params: Optional[Dict[str, Any]] = None

@api_router.post("/export")
async def export_data(request: ExportRequest):
    """Export query results in various formats"""
    try:
        data = request.data
        columns = data.get('columns', [])
        rows = data.get('rows', [])
        
        if request.format == 'csv':
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(columns)
            writer.writerows(rows)
            return {"content": output.getvalue(), "filename": f"{request.filename}.csv"}
        
        elif request.format == 'json':
            result = []
            for row in rows:
                result.append(dict(zip(columns, row)))
            return {"content": json.dumps(result, indent=2), "filename": f"{request.filename}.json"}
        
        elif request.format == 'excel':
            wb = Workbook()
            ws = wb.active
            ws.append(columns)
            for row in rows:
                ws.append(row)
            
            excel_output = io.BytesIO()
            wb.save(excel_output)
            excel_output.seek(0)
            import base64
            content = base64.b64encode(excel_output.getvalue()).decode()
            return {"content": content, "filename": f"{request.filename}.xlsx", "encoding": "base64"}
        
        elif request.format == 'sql':
            output = []
            table_name = request.filename or "exported_table"
            for row in rows:
                values = []
                for val in row:
                    if val is None:
                        values.append("NULL")
                    elif isinstance(val, str):
                        escaped_val = val.replace("'", "''")
                        values.append(f"'{escaped_val}'")
                    else:
                        values.append(str(val))
                output.append(f"INSERT INTO {table_name} ({', '.join(columns)}) VALUES ({', '.join(values)});")
            
            return {"content": "\n".join(output), "filename": f"{request.filename}.sql"}
        
        else:
            raise HTTPException(status_code=400, detail="Unsupported export format")
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@api_router.post("/format-query")
async def format_query(request: FormatQueryRequest):
    """Format SQL query for better readability"""
    try:
        formatted = sqlparse.format(
            request.query,
            reindent=True,
            keyword_case='upper',
            identifier_case='lower',
            strip_comments=False,
            use_space_around_operators=True
        )
        return {"formatted_query": formatted}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@api_router.post("/schema")
async def get_schema(request: SchemaRequest):
    """Get database schema information"""
    try:
        db_type = request.database_type.lower()
        schema = {"tables": []}
        
        if db_type == "sqlite":
            db_path = request.connection_params.get('database', '/tmp/sql_studio_default.db') if request.connection_params else '/tmp/sql_studio_default.db'
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            
            # Get all tables
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")
            tables = cursor.fetchall()
            
            for table in tables:
                table_name = table[0]
                # Get columns for each table
                cursor.execute(f"PRAGMA table_info({table_name})")
                columns = cursor.fetchall()
                
                table_info = {
                    "name": table_name,
                    "columns": [
                        {
                            "name": col[1],
                            "type": col[2],
                            "nullable": not col[3],
                            "primary_key": bool(col[5])
                        }
                        for col in columns
                    ]
                }
                schema["tables"].append(table_info)
            
            conn.close()
            return schema
        
        elif db_type == "mysql":
            if not request.connection_params:
                raise ValueError("MySQL connection parameters required")
            
            conn = pymysql.connect(
                host=request.connection_params.get('host', 'localhost'),
                port=request.connection_params.get('port', 3306),
                user=request.connection_params.get('user', 'root'),
                password=request.connection_params.get('password', ''),
                database=request.connection_params.get('database', ''),
                charset='utf8mb4'
            )
            cursor = conn.cursor()
            
            # Get all tables
            cursor.execute("SHOW TABLES")
            tables = cursor.fetchall()
            
            for table in tables:
                table_name = table[0]
                cursor.execute(f"DESCRIBE {table_name}")
                columns = cursor.fetchall()
                
                table_info = {
                    "name": table_name,
                    "columns": [
                        {
                            "name": col[0],
                            "type": col[1],
                            "nullable": col[2] == 'YES',
                            "primary_key": col[3] == 'PRI'
                        }
                        for col in columns
                    ]
                }
                schema["tables"].append(table_info)
            
            conn.close()
            return schema
        
        elif db_type == "postgresql":
            if not request.connection_params:
                raise ValueError("PostgreSQL connection parameters required")
            
            conn = psycopg2.connect(
                host=request.connection_params.get('host', 'localhost'),
                port=request.connection_params.get('port', 5432),
                user=request.connection_params.get('user', 'postgres'),
                password=request.connection_params.get('password', ''),
                database=request.connection_params.get('database', 'postgres')
            )
            cursor = conn.cursor()
            
            cursor.execute("""
                SELECT table_name 
                FROM information_schema.tables 
                WHERE table_schema = 'public' AND table_type = 'BASE TABLE'
            """)
            tables = cursor.fetchall()
            
            for table in tables:
                table_name = table[0]
                cursor.execute(f"""
                    SELECT column_name, data_type, is_nullable
                    FROM information_schema.columns 
                    WHERE table_name = '{table_name}'
                """)
                columns = cursor.fetchall()
                
                table_info = {
                    "name": table_name,
                    "columns": [
                        {
                            "name": col[0],
                            "type": col[1],
                            "nullable": col[2] == 'YES',
                            "primary_key": False
                        }
                        for col in columns
                    ]
                }
                schema["tables"].append(table_info)
            
            conn.close()
            return schema
        
        else:
            raise HTTPException(status_code=400, detail="Unsupported database type for schema exploration")
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@api_router.post("/import-data")
async def import_data(request: ImportDataRequest):
    """Import data into a table"""
    try:
        db_type = request.database_type.lower()
        
        if db_type == "sqlite":
            db_path = request.connection_params.get('database', '/tmp/sql_studio_default.db') if request.connection_params else '/tmp/sql_studio_default.db'
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            
            if request.data:
                columns = list(request.data[0].keys())
                placeholders = ','.join(['?' for _ in columns])
                
                for row in request.data:
                    values = [row.get(col) for col in columns]
                    cursor.execute(
                        f"INSERT INTO {request.table_name} ({','.join(columns)}) VALUES ({placeholders})",
                        values
                    )
                
                conn.commit()
            
            conn.close()
            return {"success": True, "rows_imported": len(request.data)}
        
        else:
            raise HTTPException(status_code=400, detail="Import currently only supported for SQLite")
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@api_router.get("/snippets")
async def get_sql_snippets():
    """Get common SQL query snippets"""
    snippets = [
        {
            "name": "Select All",
            "description": "Select all records from a table",
            "query": "SELECT * FROM table_name;"
        },
        {
            "name": "Select with WHERE",
            "description": "Select with condition",
            "query": "SELECT column1, column2\\nFROM table_name\\nWHERE condition;"
        },
        {
            "name": "Inner Join",
            "description": "Join two tables",
            "query": "SELECT t1.*, t2.*\\nFROM table1 t1\\nINNER JOIN table2 t2 ON t1.id = t2.foreign_id;"
        },
        {
            "name": "Group By with Aggregation",
            "description": "Group and aggregate data",
            "query": "SELECT column, COUNT(*), SUM(amount)\\nFROM table_name\\nGROUP BY column\\nHAVING COUNT(*) > 1;"
        },
        {
            "name": "Create Table",
            "description": "Create a new table",
            "query": "CREATE TABLE table_name (\\n    id INTEGER PRIMARY KEY,\\n    name TEXT NOT NULL,\\n    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP\\n);"
        },
        {
            "name": "Insert Data",
            "description": "Insert new records",
            "query": "INSERT INTO table_name (column1, column2)\\nVALUES ('value1', 'value2');"
        },
        {
            "name": "Update Records",
            "description": "Update existing records",
            "query": "UPDATE table_name\\nSET column1 = 'new_value'\\nWHERE condition;"
        },
        {
            "name": "Delete Records",
            "description": "Delete records",
            "query": "DELETE FROM table_name\\nWHERE condition;"
        },
        {
            "name": "Subquery",
            "description": "Query with subquery",
            "query": "SELECT *\\nFROM table_name\\nWHERE column IN (\\n    SELECT column FROM other_table WHERE condition\\n);"
        },
        {
            "name": "Transaction",
            "description": "Transaction block",
            "query": "BEGIN TRANSACTION;\\n-- Your SQL statements here\\nCOMMIT;\\n-- Or use ROLLBACK; to undo"
        }
    ]
    return {"snippets": snippets}

@api_router.get("/stats")
async def get_query_stats():
    """Get query execution statistics"""
    try:
        # Get total queries executed
        total = await db.query_history.count_documents({})
        
        # Get success rate
        successful = await db.query_history.count_documents({"success": True})
        failed = await db.query_history.count_documents({"success": False})
        
        # Get database distribution
        pipeline = [
            {"$group": {"_id": "$database_type", "count": {"$sum": 1}}},
            {"$sort": {"count": -1}}
        ]
        db_distribution = await db.query_history.aggregate(pipeline).to_list(100)
        
        # Get recent activity (last 7 days)
        from datetime import datetime, timedelta, timezone
        seven_days_ago = datetime.now(timezone.utc) - timedelta(days=7)
        recent_count = await db.query_history.count_documents({
            "timestamp": {"$gte": seven_days_ago.isoformat()}
        })
        
        return {
            "total_queries": total,
            "successful_queries": successful,
            "failed_queries": failed,
            "success_rate": round((successful / total * 100) if total > 0 else 0, 2),
            "database_distribution": [
                {"database": item["_id"], "count": item["count"]}
                for item in db_distribution
            ],
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("startup")
async def startup_db():
    """Create sample database on startup"""
    try:
        db_path = "/tmp/sql_studio_default.db"
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        # Check if tables already exist
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='users'")
        if not cursor.fetchone():
            # Create sample tables
            cursor.execute("""
                CREATE TABLE users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    email TEXT UNIQUE NOT NULL,
                    age INTEGER,
                    department TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
            
            cursor.execute("""
                CREATE TABLE orders (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER,
                    product TEXT NOT NULL,
                    amount DECIMAL(10, 2),
                    status TEXT DEFAULT 'pending',
                    order_date TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (user_id) REFERENCES users(id)
                )
            """)
            
            cursor.execute("""
                CREATE TABLE products (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    category TEXT,
                    price DECIMAL(10, 2),
                    stock INTEGER DEFAULT 0
                )
            """)
            
            # Insert sample data
            cursor.executemany(
                "INSERT INTO users (name, email, age, department) VALUES (?, ?, ?, ?)",
                [
                    ("John Doe", "john@example.com", 30, "Engineering"),
                    ("Jane Smith", "jane@example.com", 25, "Marketing"),
                    ("Bob Johnson", "bob@example.com", 35, "Sales"),
                    ("Alice Brown", "alice@example.com", 28, "Engineering"),
                    ("Charlie Davis", "charlie@example.com", 32, "HR"),
                    ("Eva Wilson", "eva@example.com", 27, "Marketing"),
                    ("Frank Miller", "frank@example.com", 40, "Sales"),
                    ("Grace Lee", "grace@example.com", 29, "Engineering")
                ]
            )
            
            cursor.executemany(
                "INSERT INTO products (name, category, price, stock) VALUES (?, ?, ?, ?)",
                [
                    ("Laptop", "Electronics", 999.99, 50),
                    ("Mouse", "Electronics", 29.99, 200),
                    ("Keyboard", "Electronics", 79.99, 150),
                    ("Monitor", "Electronics", 299.99, 75),
                    ("Headphones", "Electronics", 149.99, 100),
                    ("Desk Chair", "Furniture", 249.99, 30),
                    ("Desk Lamp", "Furniture", 49.99, 80),
                    ("Notebook", "Stationery", 5.99, 500)
                ]
            )
            
            cursor.executemany(
                "INSERT INTO orders (user_id, product, amount, status) VALUES (?, ?, ?, ?)",
                [
                    (1, "Laptop", 999.99, "completed"),
                    (1, "Mouse", 29.99, "completed"),
                    (2, "Keyboard", 79.99, "pending"),
                    (3, "Monitor", 299.99, "completed"),
                    (4, "Headphones", 149.99, "shipped"),
                    (5, "Desk Chair", 249.99, "completed"),
                    (6, "Desk Lamp", 49.99, "pending"),
                    (7, "Notebook", 5.99, "completed"),
                    (2, "Laptop", 999.99, "shipped"),
                    (3, "Headphones", 149.99, "completed")
                ]
            )
            
            conn.commit()
            logger.info(f"Created sample database at {db_path}")
        
        conn.close()
    except Exception as e:
        logger.error(f"Failed to create sample database: {e}")

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()

